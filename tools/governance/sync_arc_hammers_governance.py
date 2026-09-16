from pathlib import Path
from docx import Document
import hashlib
import zipfile

ROOT = Path.cwd()
DOCS = ROOT / 'docs'

AMEND = '''## Approved implementation amendment 2.16 - Kinetic Arc Hammers

Approved September 16, 2026 after Rebounding Arc Blade live acceptance. `docs/SLICE-5-ARC-HAMMERS-SCOPE.md` and ADR-077 govern this bounded increment. Preserve Section 15.10's five charges, minimum 0.35 race-second cadence, ballistic trajectories, exactly one terrain bounce, short post-bounce expiry, and the existing standard 0.85-second racer spin. Manny approved the complete operating contract while explicitly withholding gameplay implementation and asset/presentation development.

One ITEM press throws one Hammer. Normal ITEM throws forward; Brake/Reverse + ITEM throws backward. The first throw is immediately eligible and only a committed use spends a charge or starts cadence. Launch uses 36 m/s horizontal base speed plus 0.20 times the owner's planar launch velocity after a 10 m/s pre-scale cap, 11 m/s initial upward velocity, 24 m/s^2 downward gravity, 0.36 m contact radius, the existing `1.75 m + radius` launch offset, and 0.18 race seconds of owner arming. After arming, normal later owner self-hit is legal.

The first genuine descending contact with the actual supporting Circuit Alpha surface is the single terrain bounce. Retain 0.78 times tangential velocity and reflect the incoming surface-normal component at 0.55 restitution; asphalt, dirt, grass, boost and ramp surfaces use the same Hammer bounce physics. First guardrail contact destroys the Hammer and never counts as its terrain bounce. On the first terrain bounce, cap remaining lifetime to 0.75 race seconds; the next genuine terrain contact destroys it. A Hammer that never bounces has a 2.25 race-second hard lifetime. Resolve swept contacts chronologically, with same-time priority guardrail, then eligible racer, then terrain.

The first eligible racer hit applies the accepted 0.85-second standard spin through existing RacerEffects/camera/driver-state authority and destroys the Hammer, with no piercing, AoE, extra velocity penalty, Frost stack, or race-progress mutation. Generic immunity and Prismatic absorb and destroy the Hammer without a hostile effect. Shockwave clears an active Hammer at horizontal distance <=5 m before same-step Hammer movement, bounce, or racer contact. Each active Hammer owns one existing shared 40-object item-physics slot; pause freezes all Hammer timing/motion/presentation state; owner recovery or finish does not cancel an already-fired Hammer; restart, hub return, and disposal clear it.

The approved original procedural Hammer presentation and fixed-item/counter acceptance routes are future implementation contracts only. No model, VFX, audio, binary asset, or gameplay development is authorized by this governance checkpoint. No item probability, accepted-item tuning, racer statistic, track/checkpoint authority, AI item policy, dependency, or Slice 6 requirement changes.
'''

ADR = '''\n## ADR-077: Bound Kinetic Arc Hammers to one physical terrain rebound

- **Date:** 2026-09-16
- **Status:** Approved by Manny; governance publication checkpoint in progress.
- **Approval:** Manny approved the complete `docs/SLICE-5-ARC-HAMMERS-SCOPE.md` contract and explicitly directed that gameplay implementation and asset/presentation development remain off the table.
- **Context:** PRD Section 15.10 fixes five charges, a 0.35-second minimum cadence, ballistic trajectories, exactly one terrain bounce and short post-bounce expiry, while leaving launch physics, bounce coefficients, collision ordering, exact lifetimes, counters and lifecycle behavior unspecified. The existing runtime has reusable inventory, capacity, RacerEffects, immunity and Shockwave boundaries but no ballistic terrain-bounce projectile path.
- **Decision:** Amendment 2.16 and the approved scope define forward/backward use; 36 m/s horizontal launch, 11 m/s upward velocity, 24 m/s^2 gravity, 0.20x capped planar inheritance, 0.36 m radius, 0.18-second owner arming; one actual supporting-surface rebound with 0.78 tangential retention and 0.55 normal restitution; 0.75-second post-bounce and 2.25-second hard lifetime; first-wall destruction; standard 0.85-second hit spin and destruction; later owner self-hit; Prismatic/generic immunity absorption; Shockwave pre-movement clearing; guardrail > racer > terrain same-time ordering; and existing shared-capacity/pause/cleanup authority.
- **Architecture:** Any later authorized implementation keeps Hammer ballistic state inside the existing ProjectileSystem/item ownership boundary and samples the actual supporting race surface. It may add only the focused surface-query support necessary for the governed bounce; it does not move item logic into kart physics or redesign Circuit Alpha.
- **Rationale:** The approved numbers create a visibly lobbed projectile that remains useful at kart-racing speeds while the short post-bounce interval prevents five charges from becoming persistent track clutter. Wall destruction keeps the Hammer mechanically distinct from Ricochet Kinetic Disc.
- **Consequences:** A later separately authorized implementation requires bounded 3D ballistic/surface-contact tests, but no new gameplay subsystem, capacity pool, track rewrite, dependency, binary asset, or AI item policy. Gameplay and asset/presentation development remain separately gated and are not authorized by this decision publication.
'''

HAMMER_SCOPE = '''## Kinetic Arc Hammers scope approval - approved 2026-09-16

PRD amendment 2.16 / ADR-077 and `docs/SLICE-5-ARC-HAMMERS-SCOPE.md` govern the next bounded item contract. The approved fill-ins are forward/backward directional throws; 36 m/s horizontal launch, 11 m/s upward velocity, 24 m/s^2 gravity and 0.20x capped planar inheritance; 0.36 m radius and 0.18-second owner arming; exactly one actual supporting-surface rebound with 0.78 tangential retention and 0.55 normal restitution; 0.75-second post-bounce and 2.25-second hard lifetime; first-guardrail destruction; standard 0.85-second hit spin/destruction; later owner self-hit; Prismatic/generic immunity absorption; Shockwave <=5 m pre-movement clearing; deterministic guardrail > racer > terrain same-time priority; shared 40-object capacity; and pause/lifecycle cleanup. The future original procedural presentation and acceptance routes are approved as a contract only. **Manny explicitly withheld Arc Hammers gameplay implementation and asset/presentation development.** Keep the Arc Hammers functional checklist row open until separately authorized gameplay is implemented, validated, deployed and live accepted.

'''

ARC_TEST = '''## Slice 5 Arc Blade - deployed and LIVE ACCEPTED

`docs/SLICE-5-ARC-BLADE-SCOPE.md` and amendment 2.15 / ADR-076 govern the accepted implementation. Governance PR #138 merged at `7ce6511bc040d2b176ed528b687ed589fafd045d`; gameplay PR #139 merged at `8822341b61900799e0166cfe94bf69cb3986bf0e`. Hosted PR CI `35118244169` and post-merge validation/Pages `35118484183` passed with **48 files / 431 tests**, **82.50% statement coverage**, strict typecheck, zero-warning lint, asset verification and production build. Manny reported that all deployed Arc Blade live tests passed on September 16, 2026; PR #139 comment `5700594653` is the product-owner evidence. No browser/device versions beyond that explicit report are inferred.

| Coverage | Evidence |
| --- | --- |
| Inventory, movement and contacts | `tests/arc-blade.test.ts`: three charges and exact cadence, rollback without cues, forward/reverse equivalence, measured curved path and moving-owner return, arming/catch/expiry boundaries, per-leg hit isolation and turnaround separation, chronological wall/contact/catch ordering, immunity/counters, shared mixed-object capacity and owner cancellation. Real Circuit Alpha curves and elevated sections are included. |
| Actual runtime and input | `tests/arc-runtime.test.ts`: moving keyboard/mobile ITEM paths, ordinary race-state guards, actual Rapier/controller spin and camera/hit state, independent Frost/Nitro/Prismatic state, before-contact Shockwave clearing, recovery and all three moving diagnostics. Counter PASS requires a real measured encounter; a miss remains INCONCLUSIVE. |
| Presentation and cleanup | `tests/arc-presentation.test.ts`: finite trail/return accent, mobile/desktop chase/rear frustums after the camera intro, bounded flashes and audio voices, gesture unlock/volume/pause/unavailable audio/disposal, and 200 complete throws returning resource counts to baseline. |
| Stress | `tests/arc-soak.test.ts`: 800 throws with 40 simultaneous objects and eight racer snapshots, finite transforms and bounded capacity/resources. The isolated CPU observation was local Node/JSDOM evidence, not final rendered-device performance certification. |

### Arc Blade deployed regression routes

Use `?testItem=arc-blade` to race normally and verify the accepted three-charge/cadence, forward-only throw under normal/reverse ITEM intent, curved outbound/return path, safe owner catch, rival spin and legitimate separated second-leg hit behavior. Optional deployed counter routes remain:

- `?testItem=shockwave&testArcBladeCounter=shockwave`
- `?testItem=prismatic-invincibility&testArcBladeCounter=prismatic&testArcBladePhase=protected`
- `?testItem=prismatic-invincibility&testArcBladeCounter=prismatic&testArcBladePhase=expired`

A miss, interception, wall/catch/expiry, recovery or unsuitable geometry remains INCONCLUSIVE rather than a protection pass. Preserve the accepted normal-build regression matrix, both cameras, desktop/mobile controls, readable original audiovisuals, and lifecycle cleanup in future regression testing.

## Slice 5 Kinetic Arc Hammers - approved governance contract; gameplay/assets held

PRD amendment 2.16, ADR-077, and `docs/SLICE-5-ARC-HAMMERS-SCOPE.md` define the approved future validation contract. **Manny explicitly did not authorize Arc Hammers gameplay implementation or asset/presentation development.** Therefore no Hammer runtime route, model, VFX/audio implementation, fixture, or gameplay test is deployed by this governance checkpoint.

When gameplay is separately authorized later, automated validation must cover the exact approved values and boundaries: five charges; 0.35-second commit-only cadence; forward/backward input; 36 m/s horizontal launch, 11 m/s upward velocity, 24 m/s^2 gravity and 0.20x capped planar inheritance; 0.36 m radius; 0.18-second owner arming; one genuine supporting-surface bounce with 0.78 tangential retention / 0.55 normal restitution; 0.75-second post-bounce and 2.25-second hard lifetime; first-wall and second-terrain-contact destruction; 0.85-second standard hit spin; later self-hit; immunity/Prismatic absorption; Shockwave <=5 m pre-movement clearing; chronological collision ordering; shared-capacity rollback; pause/lifecycle cleanup; race-authority preservation; repeated-use soak; and full accepted-item/race regressions. Helper-only trajectories are insufficient: the real dispatcher, capacity, Circuit Alpha surface query, RacerEffects and runtime camera/contact path must be exercised.

Future primary deployed review route, only after separately authorized gameplay publication: `?testItem=arc-hammers`. Future optional diagnostics are `?testItem=shockwave&testArcHammerCounter=shockwave` and the Prismatic protected/expired `testArcHammerCounter=prismatic` routes defined in the scope. They must never be treated as live or distributed as playable Hammer tests before that later deployment.

'''

STATUS = '''# Implementation Status

## Current slice

**Slice 5 - Item Boxes, Weapons & Position-Based Distribution - REBOUNDING ARC BLADE LIVE ACCEPTED; KINETIC ARC HAMMERS SCOPE APPROVED / GOVERNANCE PUBLICATION PENDING**

PRD baseline on `main`: **v1.1, approved implementation amendment 2.15**. The PR #142 governance branch synchronizes approved amendment 2.16 / ADR-077 for publication.

Latest verified `main`: **`da6e7e178d30ed5e2d03103dd1bbe18ba120b97d`**.

**Hard hold:** Manny approved the complete Kinetic Arc Hammers scope on September 16, 2026 and explicitly directed that Arc Hammers gameplay development and asset/presentation development remain off the table. The current checkpoint is documentation/governance only. Gameplay requires a separate later authorization after governance publication is fully cleared.

## Rebounding Arc Blade final publication and live acceptance - 2026-09-16

Rebounding Arc Blade is **LIVE ACCEPTED** under amendment 2.15 / ADR-076. Governance PR #138 merged at `7ce6511bc040d2b176ed528b687ed589fafd045d`; gameplay PR #139 merged at `8822341b61900799e0166cfe94bf69cb3986bf0e`; hosted PR CI `35118244169` and post-merge validation/Pages `35118484183` passed with **48 test files / 431 tests**, **82.50% statement coverage**, strict typecheck, zero-warning lint, runtime/branding asset checks and production build. Manny reported all deployed Arc Blade live tests passed; PR #139 comment `5700594653` records the product-owner result. Acceptance reconciliation PR #140 and README alignment PR #141 subsequently merged, with post-merge runs `35120596434` and `35121093818` passing.

## Kinetic Arc Hammers approved scope - 2026-09-16

`docs/SLICE-5-ARC-HAMMERS-SCOPE.md`, amendment 2.16 and ADR-077 define the approved bounded contract: five charges; 0.35-second minimum commit cadence; forward/backward ballistic throws; 36 m/s horizontal launch, 11 m/s upward velocity, 24 m/s^2 gravity, 0.20x capped planar inheritance; 0.36 m radius; 0.18-second owner arming; one supporting-surface bounce with 0.78 tangential retention / 0.55 normal restitution; 0.75-second post-bounce and 2.25-second hard lifetime; first-wall and second-terrain-contact destruction; standard 0.85-second hit spin and destruction; later owner self-hit; generic/Prismatic immunity absorption; Shockwave <=5 m pre-movement clearing; guardrail > racer > terrain same-time priority; shared 40-object capacity; and pause/lifecycle cleanup.

The future original procedural presentation and acceptance routes are approved **only as contract language**. No Hammer runtime source, gameplay code, model, VFX, audio, binary asset, fixture, or playable test route is authorized or created by this checkpoint.

## Governance publication checkpoint

The PR #142 governance tree synchronizes `docs/PRD.md`, `docs/DECISIONS.md`, `docs/SLICE-5-ITEM-SYSTEM-DESIGN.md`, `docs/TESTING.md`, `docs/IMPLEMENTATION-STATUS.md`, `README.md`, `docs/SLICE-5-ARC-HAMMERS-SCOPE.md`, and the Word PRD approval artifact with amendment 2.16 / ADR-077. The Word artifact was rendered after synchronization: 52 pages, with pages 1-44 and 49-50 pixel-identical to the approved 2.15 artifact and all changed/new appendix pages visually reviewed without clipping, overlap, missing glyphs, or footer errors.

No gameplay source or runtime asset is part of this governance synchronization. Before merge, the final tree must pass full repository validation, Git LFS verification and hosted PR CI. Temporary governance transfer tooling must be removed before merge. After merge, verify post-merge validation/Pages and then **STOP**. Gameplay/assets remain held.

## Slice 5 accepted functional increments

Deployed and live accepted bounded increments include item boxes/one-slot inventory/roulette/HUD/input foundation, Nitro Surge, Ricochet Kinetic Disc, Homing Seeker Drone, Apex core, Timed Blast Orb/HazardSystem, Slick Trap, Slick/Blast AI hazard response, Acoustic Shockwave Pulse, Prismatic Invincibility, Blaze Orbs, Frost Orbs, and Rebounding Arc Blade. The seeded distribution evidence from PR #114 remains valid.

Four item effects remain unimplemented and not live accepted: **Kinetic Arc Hammers, Vision-Obscuring Ink Splat, Continuous Nitro Overdrive, and Hyper-Drive Rocket**. Arc Hammers is scope-approved but implementation-held.

## Known defects / unresolved items

- Issue #106 remains future development and nonblocking for accepted item increments.
- The existing production-build large-chunk warning remains known and nonblocking.
- Three moderate npm audit findings remain in the unchanged dependency set; no remediation is part of this governance increment.
- Full AI item acquisition/use, final all-item interaction/counter matrix, final lifecycle/object-count soak, final item/VFX performance evidence, and final desktop/mobile Slice 5 acceptance remain open.
- Slice 6 remains locked.

## Next recommended action

Complete and publish **governance documentation only** for amendment 2.16 / ADR-077. Require passing hosted CI before merge, verify post-merge validation/Pages, record the governance checkpoint, then stop. **Do not begin Kinetic Arc Hammers gameplay or asset/presentation development without a new explicit Manny authorization.**

## Approval state

**Slice 5 design:** APPROVED 2026-09-05.

**Rebounding Arc Blade:** LIVE ACCEPTED 2026-09-16.

**Kinetic Arc Hammers scope:** APPROVED 2026-09-16.

**Kinetic Arc Hammers governance publication:** PENDING PR #142 merge + post-merge verification.

**Kinetic Arc Hammers gameplay:** NOT AUTHORIZED.

**Kinetic Arc Hammers asset/presentation development:** NOT AUTHORIZED.

**Slice 5 implementation:** IN PROGRESS.

**Slice 6:** LOCKED pending full Slice 5 validation, deployment, and Manny live acceptance.
'''


def update(rel, fn):
    path = ROOT / rel
    old = path.read_text()
    new = fn(old)
    if new == old:
        raise RuntimeError(f'No change for {rel}')
    path.write_text(new)


def prd(s):
    anchor = "The primary test route will force Arc pickups while preserving unrestricted normal-race input. The scope document specifies the exact geometry, lifecycle, presentation and automated/live gates. This governance checkpoint changes no gameplay. No probability, dependency, AI item-use, other item balance or Slice 6 change is authorized. Verify governance merge and CI/Pages before the separately authorized gameplay increment; gameplay publication and live acceptance remain later gates.\n\n"
    assert anchor in s and 'Approved implementation amendment 2.16' not in s
    return s.replace(anchor, anchor + AMEND + '\n', 1)


def decisions(s):
    old = '- **Status:** Governance published through PR #138 at `7ce6511bc040d2b176ed528b687ed589fafd045d`, with CI/Pages `35027045477` successful. Manny subsequently authorized gameplay (“Proceed”) and continued it on September 16. Gameplay is locally validated; publication and live acceptance remain pending.'
    new = '- **Status:** LIVE ACCEPTED. Governance PR #138 and gameplay PR #139 are merged/deployed; post-merge gameplay validation/Pages `35118484183` passed. Manny reported all deployed Arc Blade live tests passed on September 16, 2026; PR #139 comment `5700594653` records the product-owner evidence.'
    assert old in s
    s = s.replace(old, new, 1)
    old = '- **Gate:** Governance and gameplay implementation authorization are satisfied. Local validation passed 48 files / 431 tests and all standard checks on September 16. Gameplay publication approval, hosted CI, post-merge validation/Pages and live acceptance remain open; no rendered/device result is inferred from automated tests.'
    new = '- **Gate:** Closed. Hosted PR CI `35118244169`, merge `8822341b61900799e0166cfe94bf69cb3986bf0e`, post-merge validation/Pages `35118484183`, and Manny\'s all-tests-pass live acceptance complete the Arc Blade increment. No device/browser versions beyond that explicit report are inferred.'
    assert old in s and 'ADR-077:' not in s
    return s.replace(old, new, 1).rstrip() + ADR + '\n'


def readme(s):
    old = 'Kinetic Arc Hammers is the next approval-gated scope proposal; no Hammer gameplay is authorized yet.'
    new = 'Kinetic Arc Hammers scope is approved under amendment 2.16 / ADR-077 and the current checkpoint is governance publication only; Hammer gameplay and asset/presentation development are explicitly not authorized.'
    assert old in s
    return s.replace(old, new, 1)


def item_design(s):
    assert '- [ ] Arc Blade provides three charges' in s
    s = s.replace('- [ ] Arc Blade provides three charges', '- [x] Arc Blade provides three charges', 1)
    marker = '# Slice 5 exit checklist\n'
    assert marker in s and 'Kinetic Arc Hammers scope approval - approved 2026-09-16' not in s
    s = s.replace(marker, HAMMER_SCOPE + marker, 1)
    old = "Frost acceptance reconciliation: PR #137 comment `5625755529` and `docs/SLICE-5-FROST-ORBS-SCOPE.md` at `f44176ceef38dfeebee10f6d44a0c41fdb869629` record Manny's final live acceptance. Follow-on validation/Pages `34533224332` passed. This closes the bounded Frost functional row; final all-item gates remain open. Rebounding Arc Blade is locally implemented and validated under `docs/SLICE-5-ARC-BLADE-SCOPE.md`, amendment 2.15 / ADR-076, following governance PR #138 and Manny's gameplay authorization. Local validation on September 16 passed 48 files / 431 tests and the full validation command. Gameplay publication and live acceptance remain pending, so the Arc acceptance row remains open."
    new = "Frost acceptance reconciliation: PR #137 comment `5625755529` and `docs/SLICE-5-FROST-ORBS-SCOPE.md` at `f44176ceef38dfeebee10f6d44a0c41fdb869629` record Manny's final live acceptance. Follow-on validation/Pages `34533224332` passed. Rebounding Arc Blade is also live accepted: gameplay PR #139 merged at `8822341b61900799e0166cfe94bf69cb3986bf0e`, post-merge validation/Pages `35118484183` passed, and PR #139 comment `5700594653` records Manny's all-tests-pass deployed acceptance. Arc Hammers scope is approved under amendment 2.16 / ADR-077, but gameplay and asset/presentation development remain explicitly unauthorized, so its functional row stays open. Final all-item gates remain open."
    assert old in s
    return s.replace(old, new, 1)


def testing(s):
    start = s.index('## Slice 5 Arc Blade')
    end = s.index('## Slice 5 Frost Orbs')
    return s[:start] + ARC_TEST + s[end:]

update('docs/PRD.md', prd)
update('docs/DECISIONS.md', decisions)
update('README.md', readme)
update('docs/SLICE-5-ITEM-SYSTEM-DESIGN.md', item_design)
update('docs/TESTING.md', testing)
(DOCS / 'IMPLEMENTATION-STATUS.md').write_text(STATUS)

word = DOCS / 'Manacondas_Minigame_Mayhem_PRD_v1.1.docx'
doc = Document(word)
intro = next((p for p in doc.paragraphs if p.style.name == 'Amendment Body' and p.text.startswith('This appendix records the approved working PRD amendments 2.3 through 2.15.')), None)
assert intro is not None and not any(p.text.startswith('Amendment 2.16') for p in doc.paragraphs)
intro.text = ("This appendix records the approved working PRD amendments 2.3 through 2.16. They supersede corresponding earlier baseline clauses. Amendment 2.15 records Manny's September 11, 2026 approval of Rebounding Arc Blade scope/governance, which is now live accepted. Amendment 2.16 records Manny's September 16, 2026 approval of Kinetic Arc Hammers scope while explicitly withholding gameplay and asset/presentation development. The canonical working text and referenced scope documents are maintained in docs/PRD.md in Manaconda33/manacondas-minigame-mayhem.")
doc.add_paragraph('Amendment 2.16 Kinetic Arc Hammers', style='Amendment Heading')
for text in [
    "Approved September 16, 2026 after Rebounding Arc Blade live acceptance. docs/SLICE-5-ARC-HAMMERS-SCOPE.md and ADR-077 govern this bounded increment. Preserve Section 15.10's five charges, minimum 0.35 race-second cadence, ballistic trajectories, exactly one terrain bounce, short post-bounce expiry, and the existing standard 0.85-second racer spin. Manny approved the complete operating contract while explicitly withholding gameplay implementation and asset/presentation development.",
    "One ITEM press throws one Hammer. Normal ITEM throws forward; Brake/Reverse + ITEM throws backward. The first throw is immediately eligible and only a committed use spends a charge or starts cadence. Launch uses 36 m/s horizontal base speed plus 0.20 times the owner's planar launch velocity after a 10 m/s pre-scale cap, 11 m/s initial upward velocity, 24 m/s^2 downward gravity, 0.36 m contact radius, the existing 1.75 m plus radius launch offset, and 0.18 race seconds of owner arming. After arming, normal later owner self-hit is legal.",
    "The first genuine descending contact with the actual supporting Circuit Alpha surface is the single terrain bounce. Retain 0.78 times tangential velocity and reflect the incoming surface-normal component at 0.55 restitution; asphalt, dirt, grass, boost and ramp surfaces use the same Hammer bounce physics. First guardrail contact destroys the Hammer and never counts as its terrain bounce. On the first terrain bounce, cap remaining lifetime to 0.75 race seconds; the next genuine terrain contact destroys it. A Hammer that never bounces has a 2.25 race-second hard lifetime. Resolve swept contacts chronologically, with same-time priority guardrail, then eligible racer, then terrain.",
    "The first eligible racer hit applies the accepted 0.85-second standard spin through existing RacerEffects/camera/driver-state authority and destroys the Hammer, with no piercing, AoE, extra velocity penalty, Frost stack, or race-progress mutation. Generic immunity and Prismatic absorb and destroy the Hammer without a hostile effect. Shockwave clears an active Hammer at horizontal distance less than or equal to 5 m before same-step Hammer movement, bounce, or racer contact. Each active Hammer owns one existing shared 40-object item-physics slot; pause freezes all Hammer timing/motion/presentation state; owner recovery or finish does not cancel an already-fired Hammer; restart, hub return, and disposal clear it.",
    "The approved original procedural Hammer presentation and fixed-item/counter acceptance routes are future implementation contracts only. No model, VFX, audio, binary asset, or gameplay development is authorized by this governance checkpoint. No item probability, accepted-item tuning, racer statistic, track/checkpoint authority, AI item policy, dependency, or Slice 6 requirement changes.",
]:
    doc.add_paragraph(text, style='Amendment Body')
doc.save(word)

tmp = word.with_suffix('.normalized.docx')
with zipfile.ZipFile(word, 'r') as zin, zipfile.ZipFile(tmp, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
    for name in sorted(zin.namelist()):
        info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        info.external_attr = 0o600 << 16
        zout.writestr(info, zin.read(name))
tmp.replace(word)

EXPECTED = {
    'README.md': '7e1bb33358b8e676cb8c8a322d3d3a53f2aec7c12f63e38a59f931171d22ff1a',
    'docs/PRD.md': '86c3703237d548ecfaec4add70045b4f84960e8195567a2b9f19eb3d5f443e34',
    'docs/DECISIONS.md': '9a4cd52264b5f5628ffd1acb52d595588700db01cc8ba3099ab385047ec8ed7c',
    'docs/IMPLEMENTATION-STATUS.md': 'a4d14894488433286d616181e24a3e7a61ad40394d3f59af93c5f54e0a688703',
    'docs/TESTING.md': 'c663b0fffa6a925672ab3d09836d78ea5a84180a0e240a9b688f91daee5a247c',
    'docs/SLICE-5-ITEM-SYSTEM-DESIGN.md': 'ddf1af4e800b967f88cfb3ee648d63346a9c6e20983166c35cb02c690ad6777b',
    'docs/SLICE-5-ARC-HAMMERS-SCOPE.md': 'dae2390462b5bdf451e798475d67eb74d0feed7505ca929330d9c82737ec88bd',
    'docs/Manacondas_Minigame_Mayhem_PRD_v1.1.docx': '69b9bf1495dabfda1358844aefa0a16cb5ca19b69e99adb02f095cc465d2734f',
}
for rel, expected in EXPECTED.items():
    actual = hashlib.sha256((ROOT / rel).read_bytes()).hexdigest()
    if actual != expected:
        raise RuntimeError(f'hash mismatch {rel}: {actual} != {expected}')
print('Arc Hammers governance synchronization complete and byte-verified.')
